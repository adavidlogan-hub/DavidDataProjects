"""Generate summary.docx and full_analysis.docx from counties.csv and the evidence files.

Every number in both documents is computed here from the data. Both documents open
with a definitions section that defines each term of art before it is used.

    python -m phase1.build_docs
"""
from __future__ import annotations

import collections
import csv
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from phase1.schema import BANNED_RE, ELECTIONS, TAGS  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
P1 = REPO / "phase1"
ELECTION_LABEL = {"2024_general": "Nov 5, 2024 general", "2026_primary": "Mar 3, 2026 primary",
                  "2026_runoff": "May 26, 2026 primary runoff"}
TAG_MEANING = {
    "LIVE_PRECINCT": "Precinct results updated through the night.",
    "PRECINCT_END_OF_NIGHT": "Precinct results posted on election night only once, with the final unofficial count.",
    "COUNTY_ONLY_PRECINCT_AT_CANVASS": "County totals on election night; precinct detail later (next day, at canvass, or never online).",
    "NO_SITE_OR_SOS_ONLY": "The county posts no results itself; results only through the Texas Secretary of State.",
    "UNKNOWN": "The public evidence cannot decide. Never a guess.",
}

DEFINITIONS = [
    ("Precinct", "The smallest geographic unit for which a Texas county reports votes. Some counties use vote centers or combined polling places as reporting units; where that matters it is noted."),
    ("Precinct results (precinct numbers)", "Candidate vote counts reported separately for each precinct. A table of ballots cast or turnout by precinct is not precinct results."),
    ("Election-day votes", "Votes cast in person on election day, as opposed to early in-person votes and mail (absentee) votes. In this study, precinct numbers count only if they include election-day votes."),
    ("Election night", "7:00 PM on election day to 3:00 AM the next morning, county local time (Central; Mountain for El Paso and Hudspeth counties)."),
    ("County totals", "Vote counts for the whole county, without a precinct breakdown."),
    ("Unofficial results; final unofficial count", "Results released before the canvass. The final unofficial count is the last release of election night, after all polling places have reported."),
    ("Canvass", "The official certification of results by the county, usually one to two weeks after the election. Reports marked Official or Canvass come from this stage."),
    ("Release", "A publication of new numbers. Two updates that carry identical precinct numbers are one release."),
    ("Results platform", "Where a county publishes election-night results: a vendor results site (vendor_hosted), a Clarity Elections site (clarity_style), PDF reports on the county website (county_pdfs), a site the county built itself (homegrown), or only the state site (sos_feed_only)."),
    ("Clarity Elections (Clarity, ENR)", "A vendor election night reporting service at results.enr.clarityelections.com. Each election has an election ID; every update is published as a numbered version with a timestamped detail file."),
    ("Update log", "The list of published versions a results site keeps for an election, each with its own time and data. Clarity keeps one; most other platforms do not."),
    ("Wayback Machine; capture", "The Internet Archive's public archive of web pages. A capture is a copy of a page saved at a recorded time."),
    ("Printed run time", "The date and time a tabulation report says it was produced, printed on the report or stored in the PDF file's creation date."),
    ("Upload time", "When a file was placed on the county's web server, taken from the server's Last-Modified response header. Some servers report the time of the request instead; those give no upload time."),
    ("Check 1, check 2, check 3", "The three ways this study determines timing, applied in order: (1) the vendor update log; (2) county-posted reports with printed run times and upload times; (3) Wayback Machine election-night captures."),
    ("Per-election result", "For each of the three elections studied: LIVE (precinct numbers with election-day votes in two or more distinct election-night releases), END (exactly one such release on election night), LATER (county totals only on election night; precinct numbers after election night or never online), NONE (no county-published results), or UNDETERMINED (the checks cannot decide)."),
    ("Evidence rules for LATER and END", "LATER needs a county-totals file with a printed run time or upload time on election night and no election-day precinct numbers that night; if only files dated after election night survive, the election is UNDETERMINED. END or LIVE from county-posted reports needs the precinct report's upload time on election night, or no upload time at all and a printed run time on election night; a known later upload (for example a re-upload) cannot show election-night publication."),
    ("Tag", "One label per county, computed from its three per-election results by a fixed rule: the result shared by at least two elections; a single determined election gives a low-confidence tag; otherwise UNKNOWN. Tags: " + "; ".join(f"{t}: {m}" for t, m in TAG_MEANING.items())),
    ("Confidence", "Computed, not judged: low when the tag is UNKNOWN or rests on one election; high when two or more deciding elections were settled by check 1 or check 3 content; med when two or more deciding elections include check 2 evidence or rely on the live precinct view rule."),
    ("Live precinct view rule", "An archived election-night copy of a county's own live results app that contains a precinct-by-precinct candidate results view fed by the same live data as the county totals counts as LIVE for that election, with confidence capped at med."),
    ("Verifier; disagreement rule", "An independent agent re-checks selected counties without seeing the research first. If its tag differs from the research tag, the county's final tag is UNKNOWN and both findings are logged."),
    ("SOS", "The Texas Secretary of State, whose county election officials list supplied every county's phone and email."),
    ("Bot protection; robots.txt", "Bot protection is a website feature that blocks automated requests (seen here as empty HTTP 202 or 403 responses). robots.txt is a site's published rule for automated visitors. Both were respected and never circumvented."),
]


def load():
    rows = list(csv.DictReader(open(REPO / "counties.csv", encoding="utf-8")))
    assert len(rows) == 254, f"counties.csv has {len(rows)} rows, expected 254"
    ev = {}
    for r in rows:
        ev[r["county"]] = json.loads((P1 / "evidence" / f"{r['county'].replace(' ', '_')}.json").read_text())
    ver = {}
    for p in (P1 / "verification").glob("*.json"):
        v = json.loads(p.read_text())
        if "verifier_tag" in v:
            ver[v["county"]] = v
    return rows, ev, ver


def stats(rows, ev, ver):
    s = {}
    s["n"] = len(rows)
    s["tags"] = collections.Counter(r["tag"] for r in rows)
    s["conf"] = collections.Counter((r["tag"], r["confidence"]) for r in rows)
    s["platform"] = collections.Counter(r["platform"] for r in rows)
    s["results"] = {e: collections.Counter(ev[r["county"]]["elections"][e]["result"] for r in rows) for e in ELECTIONS}
    s["checks"] = collections.Counter(x["check"] for r in rows for x in ev[r["county"]]["elections"].values()
                                      if x["result"] != "UNDETERMINED")
    s["undetermined"] = sum(1 for r in rows for x in ev[r["county"]]["elections"].values() if x["result"] == "UNDETERMINED")
    s["verified"] = len(ver)
    s["agree"] = sum(1 for v in ver.values() if v.get("agree") is True)
    s["disagree"] = sorted(c for c, v in ver.items() if v.get("agree") is False)
    s["notices"] = [r["county"] for r in rows if r["usage_notice_text"]]
    s["stated_change"] = [r["county"] for r in rows if r["stated_change_2026"] == "y"]
    return s


def doc_base(title: str, subtitle: str):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    h = d.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_paragraph(subtitle)
    return d


def add_definitions(d):
    d.add_heading("1. Definitions", level=1)
    d.add_paragraph("These terms are used throughout the document with the meanings below.")
    for term, meaning in DEFINITIONS:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(term + ". ").bold = True
        p.add_run(meaning)


def add_table(d, header, data, widths=None):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


def counts_section(d, s, num):
    d.add_heading(f"{num}. Counts by tag", level=1)
    add_table(d, ["Tag", "Counties", "high", "med", "low"],
              [[t, s["tags"].get(t, 0), s["conf"].get((t, "high"), 0), s["conf"].get((t, "med"), 0),
                s["conf"].get((t, "low"), 0)] for t in TAGS])
    d.add_paragraph()
    d.add_paragraph(f"All {s['n']} Texas counties are covered. Each row of counties.csv has the county's elections "
                    f"office phone and email from the SOS list.")


def election_section(d, s, num):
    d.add_heading(f"{num}. Results by election", level=1)
    add_table(d, ["Election", "LIVE", "END", "LATER", "NONE", "UNDETERMINED"],
              [[ELECTION_LABEL[e]] + [s["results"][e].get(k, 0) for k in ("LIVE", "END", "LATER", "NONE", "UNDETERMINED")]
               for e in ELECTIONS])
    d.add_paragraph()
    tot = sum(s["checks"].values())
    d.add_paragraph(f"Of {tot + s['undetermined']} county-elections, {tot} were determined: "
                    f"{s['checks'].get('1', 0)} by check 1 (vendor update log), {s['checks'].get('2', 0)} by check 2 "
                    f"(county-posted reports), and {s['checks'].get('3', 0)} by check 3 (Wayback captures). "
                    f"{s['undetermined']} remain UNDETERMINED.")


def method_section(d, num, short: bool):
    d.add_heading(f"{num}. Method", level=1)
    paras = [
        "Only public online sources were used. No county was contacted. Every web request went through one shared "
        "fetch module that applies per-host rate limits, honors robots.txt, retries twice, and caches every response "
        "with its time and SHA-256 hash, so raw evidence is preserved exactly.",
        "For each county and each of the three elections, the three checks were applied in order and stopped as soon as "
        "one decided the election. Check 1 reads every published update in the Clarity update log with one shared script "
        "and counts, per update, the precincts whose election-day candidate votes are nonzero. Check 2 reads the county's "
        "posted reports for their titles, printed run times, and server upload times. Check 3 reads Wayback Machine "
        "captures made between 7 PM and 3 AM local time.",
        "The county tag and its confidence are computed from the three per-election results by fixed rules, and a "
        "validator rejects any record whose tag does not follow the rule. Phone and email come from the SOS county "
        "election officials list for every county.",
        "Three rules were tightened during the run and applied to every county by dedicated passes, so results do not "
        "depend on which agent handled a county: precinct numbers count only if they include election-day votes and "
        "repeated identical numbers are one release; LATER needs positive election-night evidence; and END from "
        "county-posted reports needs an election-night upload time when one is known.",
        "Selected counties were re-checked by independent verifier agents working blind: every LIVE_PRECINCT and "
        "PRECINCT_END_OF_NIGHT, every low-confidence tag, and a seeded random 10 percent of the rest. Disagreement "
        "sets the tag to UNKNOWN.",
    ]
    if short:
        paras = paras[:2] + paras[3:5]
    for p in paras:
        d.add_paragraph(p)


def caveats_section(d, s, num):
    d.add_heading(f"{num}. Caveats", level=1)
    items = [
        "Past behavior is not a promise. Tags describe the Nov 2024, Mar 2026, and May 2026 elections; a county can change "
        f"practice for Nov 2026. Counties with a 2025 or 2026 public statement about changing precinct reporting: "
        f"{', '.join(s['stated_change']) or 'none found'}.",
        f"{s['tags'].get('UNKNOWN', 0)} counties are UNKNOWN. Most post election-night reports that are later overwritten, "
        "use web servers that do not report upload times, or block automated requests.",
        "Wayback Machine coverage of small-county results pages on election night is sparse; check 3 decided few elections. "
        "The Wayback harvest for all counties is still running and may decide more.",
        "A printed run time shows when a report was produced, not when it was posted. LATER results based on reports "
        "dated after election night cannot rule out an election-night file that was later removed.",
        "Some counties hide precinct values in a particular election (Clarity marks them protected), so the same county "
        "can be LIVE in one election and LATER in another. The two-of-three rule decides the tag.",
        "Results sites that carry a distribution or usage notice are marked internal use only: "
        f"{', '.join(s['notices']) or 'none'}.",
    ]
    for it in items:
        d.add_paragraph(it, style="List Bullet")


def lists_section(d, rows, num):
    d.add_heading(f"{num}. Counties that publish precinct results on election night", level=1)
    for tag in ("LIVE_PRECINCT", "PRECINCT_END_OF_NIGHT"):
        names = sorted(r["county"] for r in rows if r["tag"] == tag)
        conf = {r["county"]: r["confidence"] for r in rows}
        p = d.add_paragraph()
        p.add_run(f"{tag} ({len(names)}): ").bold = True
        p.add_run(", ".join(f"{n} ({conf[n]})" for n in names) or "none")


def build_summary(rows, ev, ver, s, stamp):
    d = doc_base("Texas county precinct results on election night: summary",
                 f"Phase 1 results for all 254 counties. Generated {stamp} from counties.csv.")
    add_definitions(d)
    d.add_heading("2. Headline", level=1)
    pub = s["tags"].get("LIVE_PRECINCT", 0) + s["tags"].get("PRECINCT_END_OF_NIGHT", 0)
    d.add_paragraph(
        f"{pub} counties publish precinct results on election night ({s['tags'].get('LIVE_PRECINCT', 0)} through the "
        f"night, {s['tags'].get('PRECINCT_END_OF_NIGHT', 0)} once with the final unofficial count). "
        f"{s['tags'].get('COUNTY_ONLY_PRECINCT_AT_CANVASS', 0)} publish county totals on election night and precinct "
        f"detail later. {s['tags'].get('NO_SITE_OR_SOS_ONLY', 0)} publish no results themselves. "
        f"{s['tags'].get('UNKNOWN', 0)} cannot be decided from public evidence.")
    counts_section(d, s, 3)
    lists_section(d, rows, 4)
    election_section(d, s, 5)
    method_section(d, 6, short=True)
    d.add_heading("7. Verification", level=1)
    d.add_paragraph(f"{s['verified']} counties were re-checked by a blind verifier; {s['agree']} agreed. "
                    f"Disagreements, now UNKNOWN: {', '.join(s['disagree']) or 'none'}.")
    caveats_section(d, s, 8)
    out = REPO / "summary.docx"
    d.save(out)
    return out


def build_full(rows, ev, ver, s, stamp):
    d = doc_base("Texas county precinct results on election night: full analysis",
                 f"Phase 1 results, method, verification, and per-county detail for all 254 counties. Generated {stamp}.")
    add_definitions(d)
    counts_section(d, s, 2)
    lists_section(d, rows, 3)
    election_section(d, s, 4)
    d.add_heading("5. Results platforms", level=1)
    add_table(d, ["Platform", "Counties"], sorted(s["platform"].items(), key=lambda x: -x[1]))
    d.add_paragraph()
    vend = REPO / "phase1" / "vendor"
    if (vend / "sos_voting_systems_by_county.csv").exists():
        vs = collections.Counter(r["vendor"] for r in csv.DictReader(open(vend / "sos_voting_systems_by_county.csv", encoding="utf-8")))
        d.add_paragraph("Voting system vendor by county, from the SOS Voting Systems by County list: "
                        + ", ".join(f"{k} {v}" for k, v in vs.most_common()) + ".")
    if (vend / "clarity_presence.csv").exists():
        cp = list(csv.DictReader(open(vend / "clarity_presence.csv", encoding="utf-8")))
        n200 = sum(1 for r in cp if r["elections_json_status"] == "200")
        d.add_paragraph(f"Clarity election lists: {n200} counties return a list at results.enr.clarityelections.com; "
                        "several lists are empty or incomplete, so election IDs were also taken from links on county pages "
                        "and confirmed against the county name and date inside the downloaded data.")
    method_section(d, 6, short=False)
    d.add_heading("7. Verification", level=1)
    d.add_paragraph(f"{s['verified']} counties were re-checked blind; {s['agree']} agreed with the research tag. "
                    f"Disagreements, set to UNKNOWN: {', '.join(s['disagree']) or 'none'}. "
                    "Per-county detail is in phase1/verification/ and phase1/verification_log.md.")
    caveats_section(d, s, 8)
    d.add_heading("9. Per-county results", level=1)
    d.add_paragraph("Results per election: LIVE, END, LATER, NONE, or UND (undetermined), with the check that decided it.")
    data = []
    for r in rows:
        e = ev[r["county"]]["elections"]
        cell = lambda k: ("UND" if e[k]["result"] == "UNDETERMINED" else f"{e[k]['result']} ({e[k]['check']})")
        data.append([r["county"], r["tag"], r["confidence"], cell("2024_general"), cell("2026_primary"),
                     cell("2026_runoff"), r["platform"], r["elections_office_phone"]])
    add_table(d, ["County", "Tag", "Conf", "Nov 2024", "Mar 2026", "May 2026", "Platform", "Phone"], data)
    out = REPO / "full_analysis.docx"
    d.save(out)
    return out


def check_text(path: Path) -> None:
    doc = docx.Document(str(path))
    texts = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    bad = [t for t in texts if BANNED_RE.search(t)]
    assert not bad, f"{path.name}: banned dash or emoji in {bad[:3]}"


def main() -> int:
    rows, ev, ver = load()
    s = stats(rows, ev, ver)
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    for out in (build_summary(rows, ev, ver, s, stamp), build_full(rows, ev, ver, s, stamp)):
        check_text(out)
        print(f"wrote {out.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
